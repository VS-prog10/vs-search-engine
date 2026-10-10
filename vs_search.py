# -*- coding: utf-8 -*-
"""
VS Search Engine v1.0
Автор: Владислав Сергеевич

Универсальный поисковик по файлам Excel, Word и TXT.
Поиск по 5 значениям одновременно.
"""

import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk
import os
import re
import sys
import json
import threading
import webbrowser
import subprocess
import shutil
from datetime import datetime

# --- Библиотеки для файлов ---
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_OK = True
except:
    openpyxl = None
    OPENPYXL_OK = False

try:
    from docx import Document
except:
    Document = None

# --- COM для Excel/Word ---
try:
    import win32com.client
    PYWIN32_OK = True
except:
    PYWIN32_OK = False


# ============================================================
#  ПУТИ
# ============================================================
def get_app_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))


APP_DIR = get_app_dir()
CONFIG_FILE = os.path.join(APP_DIR, "config.json")
HISTORY_FILE = os.path.join(APP_DIR, "history.json")
PINNED_FILES = os.path.join(APP_DIR, "pinned_files.json")
APP_LOG_FILE = os.path.join(APP_DIR, "app_log.txt")
BACKUP_DIR = os.path.join(APP_DIR, "backup")

MAX_HISTORY = 10


# ============================================================
#  НАСТРОЙКИ
# ============================================================
def load_config():
    default = {
        "filter_xlsx": True,
        "filter_docx": True,
        "filter_txt": True,
        "theme": "light",
        "filter_res_v1": True,
        "filter_res_v2": True,
        "filter_res_v3": True,
        "filter_res_v4": True,
        "filter_res_v5": True,
    }
    if not os.path.isfile(CONFIG_FILE):
        return default
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        for k in default:
            if k not in data:
                data[k] = default[k]
        return data
    except:
        return default


def save_config(cfg):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print("Не удалось сохранить настройки:", e)


# ============================================================
#  ЗАКРЕПЛЁННЫЕ ФАЙЛЫ
# ============================================================
def load_pinned_files():
    if not os.path.isfile(PINNED_FILES):
        return []
    try:
        with open(PINNED_FILES, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return data
        return []
    except:
        return []


def save_pinned_files(files_list):
    try:
        with open(PINNED_FILES, "w", encoding="utf-8") as f:
            json.dump(files_list, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print("Не удалось сохранить закреплённые файлы:", e)


# ============================================================
#  ИСТОРИЯ
# ============================================================
def _empty_history():
    return {"v1": [], "v2": [], "v3": [], "v4": [], "v5": []}


def load_history():
    if not os.path.isfile(HISTORY_FILE):
        return _empty_history()
    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        if "freq" in data or "callsign" in data or "state" in data:
            return _empty_history()
        if "v1" not in data:
            return _empty_history()
        return data
    except:
        return _empty_history()


def save_history(data):
    try:
        with open(HISTORY_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print("Не удалось сохранить историю:", e)


def add_to_history(key, value):
    value = value.strip()
    if not value:
        return
    hist = load_history()
    if key not in hist:
        hist[key] = []
    lst = hist[key]
    if value in lst:
        lst.remove(value)
    lst.insert(0, value)
    lst = lst[:MAX_HISTORY]
    hist[key] = lst
    save_history(hist)


# ============================================================
#  РЕЗЕРВНЫЕ КОПИИ
# ============================================================
def create_backup():
    try:
        if not os.path.isdir(BACKUP_DIR):
            os.makedirs(BACKUP_DIR)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        files_to_backup = [
            (CONFIG_FILE, f"config_{ts}.json"),
            (HISTORY_FILE, f"history_{ts}.json"),
        ]
        copied = []
        for src, dst_name in files_to_backup:
            if os.path.isfile(src):
                dst = os.path.join(BACKUP_DIR, dst_name)
                shutil.copy2(src, dst)
                copied.append(dst_name)
        if copied:
            app_log(f"Создана резервная копия: {', '.join(copied)}")
        cleanup_old_backups()
    except Exception as e:
        app_log(f"ОШИБКА резервной копии: {e}")


def cleanup_old_backups():
    try:
        if not os.path.isdir(BACKUP_DIR):
            return
        files = os.listdir(BACKUP_DIR)
        for prefix in ("config_", "history_"):
            matching = sorted([f for f in files if f.startswith(prefix)],
                              reverse=True)
            for f in matching[10:]:
                try:
                    os.remove(os.path.join(BACKUP_DIR, f))
                except:
                    pass
    except:
        pass


# ============================================================
#  ЛОГ
# ============================================================
def app_log(text):
    try:
        with open(APP_LOG_FILE, "a", encoding="utf-8") as f:
            ts = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
            f.write(f"[{ts}] {text}\n")
    except:
        pass


def app_log_start():
    try:
        with open(APP_LOG_FILE, "a", encoding="utf-8") as f:
            f.write("\n" + "=" * 60 + "\n")
            f.write("=== ЗАПУСК VS Search Engine v1.0 ===\n")
            f.write(f"Дата/время: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}\n")
            f.write(f"Пользователь Windows: {os.environ.get('USERNAME', '')}\n")
            f.write(f"Компьютер: {os.environ.get('COMPUTERNAME', '')}\n")
            f.write("=" * 60 + "\n")
    except:
        pass


def app_log_end():
    try:
        with open(APP_LOG_FILE, "a", encoding="utf-8") as f:
            f.write("=== ЗАКРЫТИЕ ПРОГРАММЫ ===\n")
            f.write(f"Дата/время: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}\n")
            f.write("=" * 60 + "\n\n")
    except:
        pass


# ============================================================
#  ГЛОБАЛЬНЫЕ
# ============================================================
selected_files = []
search_running = False
stop_requested = False
last_results = []
current_theme = "light"
v1_text_global = ""
v2_text_global = ""
v3_text_global = ""
v4_text_global = ""
v5_text_global = ""
# ============================================================
#  РУКОВОДСТВО
# ============================================================
GUIDE_TEXT = """РУКОВОДСТВО ПО ЭКСПЛУАТАЦИИ
VS Search Engine

Версия: 1.0
Автор: Владислав Сергеевич

====================================================
1. НАЗНАЧЕНИЕ
====================================================
Универсальный поисковик по локальным файлам.
Ищет до 5 значений ОДНОВРЕМЕННО в файлах:
    - Excel (.xlsx)
    - Word (.docx)
    - Текст (.txt)

====================================================
2. ЗАПУСК ПРОГРАММЫ
====================================================
Запустите файл «VSSearchEngine.exe»
двойным щелчком мыши.

При первом запуске может появиться окно
«Windows защитила ваш компьютер» - нажмите
«Подробнее» -> «Выполнить в любом случае».

====================================================
3. РАБОТА С ФАЙЛАМИ
====================================================
Кнопка «➕ Добавить файлы» - выбрать один
или несколько файлов (Ctrl или Shift).

Кнопки:
    «➖ Удалить выбранный» - удалить файл.
    «🗑 Очистить список» - удалить все.

ФИЛЬТР ФОРМАТОВ:
    Excel (.xlsx)
    Word (.docx)
    Текст (.txt)

Список файлов АВТОМАТИЧЕСКИ сохраняется
при закрытии и восстанавливается при запуске.

====================================================
4. ПОИСК ПО ЗНАЧЕНИЯМ
====================================================
Программа ищет ОДНОВРЕМЕННО до 5 значений.

Значение 1, Значение 2, Значение 3,
Значение 4, Значение 5.

Можно заполнить одно, несколько или все пять.
Пустые поля не участвуют в поиске.

Поиск идёт по ЧАСТИЧНОМУ совпадению:
    - ищем «8988» - найдётся в «8988», «8988.5», «18988»;
    - ищем «Петя» - найдётся в «Петя93883».

Каждое поле имеет историю (до 10 последних
значений) - нажмите стрелку справа.

====================================================
5. ПРОГРЕСС С ПРОЦЕНТАМИ
====================================================
Во время поиска рядом с прогресс-баром
показывается процент выполнения.

====================================================
6. ПОДСВЕТКА РЕЗУЛЬТАТОВ
====================================================
Найденные фрагменты выделяются ЗЕЛЁНЫМ ФОНОМ
с ЖИРНЫМ шрифтом.

====================================================
7. ОТКРЫТИЕ ФАЙЛОВ
====================================================
Кнопка «📂 Открыть» рядом с файлом:
    - EXCEL: файл + выделение ячейки.
    - WORD: файл + выделение абзаца.
    - TXT: окно с подсветкой и F3.

Кнопка «📁 Папка» - открывает проводник.

Файлы без совпадений НЕ открываются.

====================================================
8. ФИЛЬТР РЕЗУЛЬТАТОВ
====================================================
Галочки над результатами:
    - Значение 1
    - Значение 2
    - Значение 3
    - Значение 4
    - Значение 5

Показывают/скрывают результаты по типам.

====================================================
9. ЭКСПОРТ ОТЧЁТА
====================================================
Кнопки: «💾 Экспорт TXT», «📊 Экспорт XLSX».

Отчёт в XLSX содержит колонки:
Источник, Путь, Место, Тип, Текст, Дата.

====================================================
10. ТЁМНАЯ ТЕМА
====================================================
Кнопка «🎨 Тема» или Ctrl+D.
Тема сохраняется при следующем запуске.

====================================================
11. ГОРЯЧИЕ КЛАВИШИ
====================================================
Ctrl+Enter  - Запустить поиск
Ctrl+O      - Добавить файлы
Ctrl+L      - Очистить список файлов
Ctrl+E      - Экспорт отчёта (.txt)
Ctrl+Shift+E- Экспорт отчёта (.xlsx)
Ctrl+D      - Переключить тему
Ctrl+H      - Очистить историю
F1          - Руководство
F5          - Очистить результаты
Esc         - Остановить поиск

====================================================
12. ЛОГИ
====================================================
app_log.txt - общий лог приложения.

====================================================
13. РЕЗЕРВНЫЕ КОПИИ
====================================================
При закрытии программы копия настроек
и истории сохраняется в папку backup/.
Хранятся 10 последних копий.

====================================================
14. ВОЗМОЖНЫЕ ПРОБЛЕМЫ
====================================================
- Excel/Word не открываются с выделением -
  нужен установленный Microsoft Office.
- Антивирус блокирует - добавьте exe в исключения.
- Ошибка VCRUNTIME140.dll - установите
  VC++ 2015 Redistributable (x86).

====================================================
15. АВТОР
====================================================
Разработчик: Владислав Сергеевич
Версия: 1.0
"""


def show_guide():
    guide_win = tk.Toplevel(root)
    guide_win.title("Руководство по эксплуатации")
    guide_win.geometry("900x720")
    guide_win.transient(root)
    apply_window_theme(guide_win)

    top = tk.Frame(guide_win, bg=COLORS["bg"])
    top.pack(fill=tk.X, padx=10, pady=6)
    tk.Label(top, text="Руководство по эксплуатации",
             font=("Arial", 12, "bold"),
             bg=COLORS["bg"], fg=COLORS["fg"]).pack(side=tk.LEFT)
    tk.Button(top, text="Закрыть", command=guide_win.destroy,
              width=12).pack(side=tk.RIGHT)

    area = scrolledtext.ScrolledText(guide_win, wrap=tk.WORD,
                                     font=("Consolas", 10),
                                     bg=COLORS["entry_bg"],
                                     fg=COLORS["fg"],
                                     insertbackground=COLORS["fg"])
    area.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)
    area.insert(tk.END, GUIDE_TEXT)
    area.config(state=tk.DISABLED)

    bot = tk.Frame(guide_win, bg=COLORS["bg"])
    bot.pack(fill=tk.X, padx=10, pady=6)
    tk.Button(bot, text="Закрыть", command=guide_win.destroy,
              width=12).pack(side=tk.RIGHT)


# ============================================================
#  ФАЙЛЫ
# ============================================================
def add_files():
    files = filedialog.askopenfilenames(
        title="Выберите файлы для поиска",
        filetypes=[
            ("Поддерживаемые файлы", "*.xlsx *.docx *.txt"),
            ("Excel", "*.xlsx"),
            ("Word", "*.docx"),
            ("Текстовые", "*.txt"),
            ("Все файлы", "*.*"),
        ]
    )
    if files:
        for f in files:
            if f not in selected_files:
                selected_files.append(f)
        update_list()
        save_pinned_files(selected_files)
        app_log(f"Добавлены файлы ({len(files)} шт.)")


def remove_selected():
    selection = listbox.curselection()
    if not selection:
        messagebox.showinfo("Подсказка", "Выберите файл в списке, чтобы удалить")
        return
    for idx in reversed(selection):
        del selected_files[idx]
    update_list()
    save_pinned_files(selected_files)


def clear_list():
    if not selected_files:
        return
    if messagebox.askyesno("Подтверждение", "Очистить весь список файлов?"):
        selected_files.clear()
        update_list()
        save_pinned_files(selected_files)


def update_list():
    listbox.delete(0, tk.END)
    for f in selected_files:
        listbox.insert(tk.END, f)
    update_status_bar()


def restore_pinned_files():
    pinned = load_pinned_files()
    restored = 0
    for f in pinned:
        if os.path.isfile(f) and f not in selected_files:
            selected_files.append(f)
            restored += 1
    if restored:
        update_list()
        app_log(f"Восстановлено закреплённых файлов: {restored}")


# ============================================================
#  ЧТЕНИЕ ФАЙЛОВ
# ============================================================
def get_excel_content(filepath):
    results = []
    if openpyxl is None:
        return results
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        val = str(cell.value)
                        results.append((val,
                            f"Лист[{sheet_name}] ячейка {cell.coordinate}",
                            cell.coordinate, sheet_name))
        wb.close()
    except Exception as e:
        results.append((f"[Ошибка чтения] {str(e)}", "Ошибка", "", ""))
    return results


def get_docx_content(filepath):
    results = []
    if Document is None:
        return results
    try:
        doc = Document(filepath)
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                results.append((para.text, f"Абзац №{i+1}",
                                "paragraph", i))
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    if cell.text.strip():
                        results.append((cell.text,
                            f"Таблица {t_idx+1}, строка {r_idx+1}, столбец {c_idx+1}",
                            "table", t_idx))
    except Exception as e:
        results.append((f"[Ошибка чтения] {str(e)}", "Ошибка", "", ""))
    return results


def get_txt_content(filepath):
    results = []
    full_text = ""
    try:
        for enc in ("utf-8", "cp1251", "utf-16"):
            try:
                with open(filepath, "r", encoding=enc) as f:
                    full_text = f.read()
                break
            except UnicodeDecodeError:
                continue
    except Exception as e:
        results.append((f"[Ошибка чтения] {str(e)}", "Ошибка", 0))
        return results, ""

    lines = full_text.split("\n")
    for i, line in enumerate(lines):
        if line.strip():
            results.append((line.strip(), f"Строка №{i+1}", i+1))

    return results, full_text


# ============================================================
#  ПРОВЕРКА СОВПАДЕНИЙ (5 ЗНАЧЕНИЙ)
# ============================================================
def check_match_5(content, v1, v2, v3, v4, v5):
    types = []
    positions = []

    values = [
        (v1, "Значение 1"),
        (v2, "Значение 2"),
        (v3, "Значение 3"),
        (v4, "Значение 4"),
        (v5, "Значение 5"),
    ]

    for val, label in values:
        if not val:
            continue
        found = False
        for m in re.finditer(re.escape(val), content, re.IGNORECASE):
            positions.append((m.start(), m.end(), label))
            found = True
        if found:
            types.append(label)

    positions.sort(key=lambda x: x[0])
    return types, positions


# ============================================================
#  ОТКРЫТИЕ ФАЙЛОВ И ПАПОК
# ============================================================
def open_excel_at(filepath, sheet_name, cell_coord):
    if not PYWIN32_OK:
        messagebox.showerror("Ошибка",
            "Библиотека pywin32 не установлена.\n"
            "Выполните: pip install pywin32")
        return
    try:
        try:
            excel = win32com.client.GetActiveObject("Excel.Application")
        except:
            excel = win32com.client.Dispatch("Excel.Application")

        target_wb = None
        abs_path = os.path.abspath(filepath).lower()
        try:
            for wb in excel.Workbooks:
                try:
                    if os.path.abspath(wb.FullName).lower() == abs_path:
                        target_wb = wb
                        break
                except:
                    continue
        except:
            pass

        if target_wb is None:
            target_wb = excel.Workbooks.Open(os.path.abspath(filepath))

        excel.Visible = True
        target_wb.Activate()

        try:
            ws = target_wb.Worksheets(sheet_name)
        except:
            ws = target_wb.Worksheets(1)
        ws.Activate()

        try:
            ws.Range(cell_coord).Select()
        except:
            pass

        try:
            excel.WindowState = 1
        except:
            pass
    except Exception as e:
        messagebox.showerror("Ошибка открытия Excel",
            f"Не удалось открыть файл:\n{filepath}\n\n{e}")


def open_word_at(filepath, description):
    if not PYWIN32_OK:
        messagebox.showerror("Ошибка",
            "Библиотека pywin32 не установлена.\n"
            "Выполните: pip install pywin32")
        return
    try:
        try:
            word = win32com.client.GetActiveObject("Word.Application")
        except:
            word = win32com.client.Dispatch("Word.Application")

        target_doc = None
        abs_path = os.path.abspath(filepath).lower()
        try:
            for doc in word.Documents:
                try:
                    if os.path.abspath(doc.FullName).lower() == abs_path:
                        target_doc = doc
                        break
                except:
                    continue
        except:
            pass

        if target_doc is None:
            target_doc = word.Documents.Open(os.path.abspath(filepath))

        word.Visible = True
        target_doc.Activate()

        m = re.search(r"Абзац №(\d+)", description)
        if m:
            try:
                idx = int(m.group(1))
                if 1 <= idx <= target_doc.Paragraphs.Count:
                    para = target_doc.Paragraphs(idx)
                    para.Range.Select()
            except:
                pass
    except Exception as e:
        messagebox.showerror("Ошибка открытия Word",
            f"Не удалось открыть файл:\n{filepath}\n\n{e}")


def open_file_generic(filepath):
    try:
        os.startfile(filepath)
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{e}")


def show_in_folder(filepath):
    try:
        if not os.path.isfile(filepath):
            messagebox.showwarning("Файл не найден",
                f"Файл больше не существует:\n{filepath}")
            return
        subprocess.Popen(f'explorer /select,"{os.path.abspath(filepath)}"')
    except Exception as e:
        messagebox.showerror("Ошибка",
            f"Не удалось открыть папку:\n{e}")


# ============================================================
#  ОКНО ПРОСМОТРА TXT
# ============================================================
class TxtViewerWindow:
    def __init__(self, parent, filepath, matches):
        self.filepath = filepath
        self.matches = matches
        self.current_idx = 0

        self.win = tk.Toplevel(parent)
        self.win.title(f"Просмотр: {os.path.basename(filepath)}")
        self.win.geometry("900x700")
        apply_window_theme(self.win)

        top = tk.Frame(self.win, bg=COLORS["bg"])
        top.pack(fill=tk.X, padx=10, pady=6)

        tk.Label(top, text=f"Файл: {filepath}",
                 font=("Arial", 10, "bold"),
                 bg=COLORS["bg"], fg=COLORS["fg"]).pack(side=tk.LEFT)

        tk.Button(top, text="Закрыть",
                  command=self.win.destroy, width=10).pack(side=tk.RIGHT)

        nav = tk.Frame(self.win, bg=COLORS["bg"])
        nav.pack(fill=tk.X, padx=10, pady=4)

        tk.Label(nav, text="Переключение (F3):",
                 font=("Arial", 9),
                 bg=COLORS["bg"], fg=COLORS["fg"]).pack(side=tk.LEFT, padx=(0, 8))

        self.counter_label = tk.Label(nav, text="", font=("Arial", 9, "bold"),
                                       bg=COLORS["bg"], fg=COLORS["fg"])
        self.counter_label.pack(side=tk.LEFT, padx=(0, 10))

        tk.Button(nav, text="◀ Предыдущее (Shift+F3)",
                  command=self.prev_match, width=22).pack(side=tk.LEFT, padx=2)
        tk.Button(nav, text="Следующее (F3) ▶",
                  command=self.next_match, width=22).pack(side=tk.LEFT, padx=2)

        self.text = scrolledtext.ScrolledText(self.win, wrap=tk.WORD,
                                              font=("Consolas", 10),
                                              bg=COLORS["entry_bg"],
                                              fg=COLORS["fg"],
                                              insertbackground=COLORS["fg"])
        self.text.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

        content = ""
        for enc in ("utf-8", "cp1251", "utf-16"):
            try:
                with open(filepath, "r", encoding=enc) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue
            except:
                break

        self.text.insert("1.0", content)

        self.text.tag_config("highlight", background="yellow",
                             foreground="black")
        self.text.tag_config("current", background="#90EE90",
                             foreground="black", font=("Consolas", 10, "bold"))

        for start, end in self.matches:
            try:
                start_idx = f"1.0 + {start} chars"
                end_idx = f"1.0 + {end} chars"
                self.text.tag_add("highlight", start_idx, end_idx)
            except:
                pass

        self._update_counter()

        if self.matches:
            self._goto_match(0)

        self.win.bind("<F3>", lambda e: self.next_match())
        self.win.bind("<Shift-F3>", lambda e: self.prev_match())
        self.win.bind("<Escape>", lambda e: self.win.destroy())

        self.win.focus_force()

    def _update_counter(self):
        if self.matches:
            self.counter_label.config(
                text=f"{self.current_idx + 1} из {len(self.matches)}")
        else:
            self.counter_label.config(text="совпадений нет")

    def _goto_match(self, idx):
        if not self.matches:
            return
        idx = max(0, min(idx, len(self.matches) - 1))
        self.text.tag_remove("current", "1.0", tk.END)
        start, end = self.matches[idx]
        try:
            start_idx = f"1.0 + {start} chars"
            end_idx = f"1.0 + {end} chars"
            self.text.tag_add("current", start_idx, end_idx)
            self.text.see(start_idx)
            self.text.mark_set(tk.INSERT, start_idx)
        except:
            pass
        self.current_idx = idx
        self._update_counter()

    def next_match(self):
        if not self.matches:
            return
        new_idx = (self.current_idx + 1) % len(self.matches)
        self._goto_match(new_idx)

    def prev_match(self):
        if not self.matches:
            return
        new_idx = (self.current_idx - 1) % len(self.matches)
        self._goto_match(new_idx)
# ============================================================
#  ЦВЕТА ТЕМ
# ============================================================
COLORS = {}


def set_colors(theme):
    global COLORS
    if theme == "dark":
        COLORS = {
            "bg": "#2b2b2b",
            "fg": "#e8e8e8",
            "entry_bg": "#3c3c3c",
            "entry_fg": "#ffffff",
            "btn_bg": "#4a4a4a",
            "btn_fg": "#ffffff",
            "frame_bg": "#333333",
            "result_bg": "#1e1e1e",
            "result_fg": "#e8e8e8",
            "label_bg": "#2b2b2b",
            "label_fg": "#e8e8e8",
        }
    else:
        COLORS = {
            "bg": "#f0f0f0",
            "fg": "#000000",
            "entry_bg": "#ffffff",
            "entry_fg": "#000000",
            "btn_bg": "#e1e1e1",
            "btn_fg": "#000000",
            "frame_bg": "#f0f0f0",
            "result_bg": "#ffffff",
            "result_fg": "#000000",
            "label_bg": "#f0f0f0",
            "label_fg": "#000000",
        }


def apply_window_theme(window):
    try:
        window.configure(bg=COLORS["bg"])
    except:
        pass
    _apply_recursive(window)


def _apply_recursive(widget):
    try:
        wclass = widget.winfo_class()
    except:
        return

    try:
        if wclass in ("Frame", "Labelframe", "Toplevel"):
            widget.configure(bg=COLORS["frame_bg"])
        elif wclass == "Label":
            widget.configure(bg=COLORS["label_bg"], fg=COLORS["label_fg"])
        elif wclass == "Entry":
            widget.configure(bg=COLORS["entry_bg"], fg=COLORS["entry_fg"],
                             insertbackground=COLORS["fg"])
        elif wclass == "Text":
            widget.configure(bg=COLORS["result_bg"], fg=COLORS["result_fg"],
                             insertbackground=COLORS["fg"])
        elif wclass == "Listbox":
            widget.configure(bg=COLORS["entry_bg"], fg=COLORS["entry_fg"])
        elif wclass == "Button":
            cur_bg = str(widget.cget("bg"))
            custom_colors = ("#4CAF50", "#2196F3", "#f44336", "#FFC107")
            if cur_bg not in custom_colors:
                widget.configure(bg=COLORS["btn_bg"], fg=COLORS["btn_fg"])
        elif wclass == "Checkbutton":
            widget.configure(bg=COLORS["bg"], fg=COLORS["fg"],
                             activebackground=COLORS["bg"],
                             selectcolor=COLORS["entry_bg"])
        elif wclass == "Radiobutton":
            widget.configure(bg=COLORS["bg"], fg=COLORS["fg"])
    except:
        pass

    for child in widget.winfo_children():
        _apply_recursive(child)


def toggle_theme():
    global current_theme
    if current_theme == "light":
        current_theme = "dark"
    else:
        current_theme = "light"
    set_colors(current_theme)
    apply_window_theme(root)
    update_status_bar()
    app_log(f"Тема переключена на: {current_theme}")


def apply_current_theme():
    set_colors(current_theme)
    apply_window_theme(root)


# ============================================================
#  СТАТУС-БАР И ПРОГРЕСС
# ============================================================
def update_status_bar(text=None):
    if text is None:
        parts = []
        if search_running:
            parts.append("Поиск...")
        else:
            parts.append("Готов")
        if last_results:
            parts.append(f"Найдено: {len(last_results)}")
        if selected_files:
            parts.append(f"Файлов: {len(selected_files)}")
        status_text = "  |  ".join(parts)
    else:
        status_text = text

    try:
        status_bar_label.config(text=status_text)
    except:
        pass


def update_progress(percent, text=None):
    try:
        progress_bar["value"] = percent
        if text is None:
            text = ""
        progress_label.config(text=f"{text} {int(percent)}%")
    except:
        pass
# ============================================================
#  ЭКСПОРТ
# ============================================================
def export_txt():
    if not last_results:
        messagebox.showinfo("Нечего экспортировать",
            "Сначала выполните поиск.")
        return
    path = filedialog.asksaveasfilename(
        title="Сохранить отчёт (.txt)",
        defaultextension=".txt",
        filetypes=[("Текстовый файл", "*.txt")],
        initialfile=f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    )
    if not path:
        return
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write("ОТЧЁТ О ПОИСКЕ\n")
            f.write(f"VS Search Engine v1.0\n")
            f.write(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}\n")
            f.write("=" * 60 + "\n\n")
            for r in last_results:
                f.write(f"Источник: {r['source']}\n")
                f.write(f"Путь: {r['path']}\n")
                f.write(f"Место: {r['location']}\n")
                f.write(f"Тип: {r['types']}\n")
                f.write(f"Текст: {r['content']}\n")
                f.write("-" * 50 + "\n")
            f.write(f"\nВсего совпадений: {len(last_results)}\n")
        app_log(f"Экспорт в TXT: {path}")
        messagebox.showinfo("Готово", f"Отчёт сохранён:\n{path}")
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось сохранить: {e}")


def export_xlsx():
    if not last_results:
        messagebox.showinfo("Нечего экспортировать",
            "Сначала выполните поиск.")
        return
    if not OPENPYXL_OK:
        messagebox.showerror("Ошибка",
            "Библиотека openpyxl не установлена.\n"
            "Выполните: pip install openpyxl")
        return

    path = filedialog.asksaveasfilename(
        title="Сохранить отчёт (.xlsx)",
        defaultextension=".xlsx",
        filetypes=[("Excel-файл", "*.xlsx")],
        initialfile=f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    if not path:
        return
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Результаты поиска"

        ws["A1"] = "ОТЧЁТ О ПОИСКЕ"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "VS Search Engine v1.0"
        ws["A2"].font = Font(italic=True)
        ws["A3"] = f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
        ws["A3"].font = Font(italic=True)

        headers = ["Источник", "Путь", "Место", "Тип", "Текст", "Дата"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4CAF50",
                                    end_color="4CAF50", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        date_str = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
        for i, r in enumerate(last_results, 6):
            ws.cell(row=i, column=1, value=r.get("source", ""))
            ws.cell(row=i, column=2, value=r.get("path", ""))
            ws.cell(row=i, column=3, value=r.get("location", ""))
            ws.cell(row=i, column=4, value=r.get("types", ""))
            ws.cell(row=i, column=5, value=r.get("content", ""))
            ws.cell(row=i, column=6, value=date_str)

        total_row = len(last_results) + 7
        ws.cell(row=total_row, column=1, value="Всего совпадений:")
        ws.cell(row=total_row, column=1).font = Font(bold=True)
        ws.cell(row=total_row, column=2, value=len(last_results))
        ws.cell(row=total_row, column=2).font = Font(bold=True)

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 25
        ws.column_dimensions["E"].width = 60
        ws.column_dimensions["F"].width = 20

        wb.save(path)
        app_log(f"Экспорт в XLSX: {path}")
        messagebox.showinfo("Готово", f"Отчёт сохранён:\n{path}")
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось сохранить: {e}")


# ============================================================
#  ИСТОРИЯ
# ============================================================
def reload_history_combos():
    hist = load_history()
    v1_combo["values"] = hist.get("v1", [])
    v2_combo["values"] = hist.get("v2", [])
    v3_combo["values"] = hist.get("v3", [])
    v4_combo["values"] = hist.get("v4", [])
    v5_combo["values"] = hist.get("v5", [])


def clear_history():
    if not messagebox.askyesno("Подтверждение",
        "Очистить всю историю запросов?"):
        return
    save_history(_empty_history())
    reload_history_combos()
    app_log("История запросов очищена")
    messagebox.showinfo("Готово", "История очищена.")


# ============================================================
#  СОХРАНЕНИЕ ПРИ ВЫХОДЕ
# ============================================================
def on_close():
    cfg = {
        "filter_xlsx": bool(filter_xlsx_var.get()),
        "filter_docx": bool(filter_docx_var.get()),
        "filter_txt": bool(filter_txt_var.get()),
        "theme": current_theme,
        "filter_res_v1": bool(filter_res_v1_var.get()),
        "filter_res_v2": bool(filter_res_v2_var.get()),
        "filter_res_v3": bool(filter_res_v3_var.get()),
        "filter_res_v4": bool(filter_res_v4_var.get()),
        "filter_res_v5": bool(filter_res_v5_var.get()),
    }
    save_config(cfg)
    save_pinned_files(selected_files)
    create_backup()
    app_log("Программа закрыта пользователем")
    app_log_end()
    root.destroy()


# ============================================================
#  КОНТЕКСТНОЕ МЕНЮ
# ============================================================
def add_context_menu(widget):
    menu = tk.Menu(widget, tearoff=0)

    def copy_text():
        try:
            widget.event_generate("<<Copy>>")
        except:
            pass

    def paste_text():
        try:
            widget.event_generate("<<Paste>>")
        except:
            pass

    def cut_text():
        try:
            widget.event_generate("<<Cut>>")
        except:
            pass

    def select_all():
        try:
            if isinstance(widget, (tk.Entry, ttk.Entry, ttk.Combobox)):
                widget.select_range(0, tk.END)
            else:
                widget.tag_add(tk.SEL, "1.0", tk.END)
                widget.mark_set(tk.INSERT, "1.0")
        except:
            pass

    menu.add_command(label="Копировать", command=copy_text)
    menu.add_command(label="Вставить", command=paste_text)
    menu.add_command(label="Вырезать", command=cut_text)
    menu.add_separator()
    menu.add_command(label="Выделить всё", command=select_all)

    def show_menu(event):
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    widget.bind("<Button-3>", show_menu)


def add_result_context_menu(widget):
    menu = tk.Menu(widget, tearoff=0)

    def copy_sel():
        try:
            widget.event_generate("<<Copy>>")
        except:
            pass

    def select_all():
        try:
            widget.tag_add(tk.SEL, "1.0", tk.END)
            widget.mark_set(tk.INSERT, "1.0")
        except:
            pass

    def clear_results():
        widget.delete(1.0, tk.END)
        update_status_bar()

    menu.add_command(label="Копировать", command=copy_sel)
    menu.add_separator()
    menu.add_command(label="Выделить всё", command=select_all)
    menu.add_command(label="Очистить результаты", command=clear_results)

    def show_menu(event):
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    widget.bind("<Button-3>", show_menu)


# ============================================================
#  ПОИСК В ПОТОКЕ
# ============================================================
def search_thread_target(v1, v2, v3, v4, v5,
                         filter_xlsx, filter_docx, filter_txt):
    global search_running, stop_requested, last_results
    total_hits = 0
    last_results = []
    app_log(f"Начат поиск. Значения: '{v1}', '{v2}', '{v3}', '{v4}', '{v5}'")

    try:
        if not selected_files:
            result_text.insert(tk.END, "Нет файлов для поиска.\n", "yellow")
            return

        result_text.insert(tk.END, "─── ПОИСК ПО ФАЙЛАМ ───\n")
        result_text.insert(tk.END,
            f"Файлов для поиска: {len(selected_files)}\n")
        result_text.insert(tk.END, "=" * 60 + "\n")

        total = len(selected_files)
        for i, filepath in enumerate(selected_files, 1):
            if stop_requested:
                result_text.insert(tk.END,
                    "\n⏹ Поиск остановлен пользователем.\n", "yellow")
                break

            fname = os.path.basename(filepath)
            percent = (i / total) * 100
            update_progress(percent, f"Файлы: {i} из {total}.")
            update_status_bar(f"Поиск: {i} из {total}")
            root.update_idletasks()

            if not os.path.isfile(filepath):
                result_text.insert(tk.END,
                    f"\n⚠️ Файл не найден: {filepath}\n", "yellow")
                continue

            ext = os.path.splitext(fname)[1].lower()
            if ext == '.xlsx' and not filter_xlsx: continue
            if ext == '.docx' and not filter_docx: continue
            if ext == '.txt' and not filter_txt: continue

            if ext == '.xlsx':
                entries = get_excel_content(filepath)
                full_text = ""
            elif ext == '.docx':
                entries = get_docx_content(filepath)
                full_text = ""
            elif ext == '.txt':
                entries, full_text = get_txt_content(filepath)
            else:
                continue

            file_hits = []
            for item in entries:
                content = item[0]
                location = item[1]
                extra = item[2] if len(item) > 2 else None
                extra2 = item[3] if len(item) > 3 else None

                types, positions = check_match_5(content, v1, v2, v3, v4, v5)
                if types:
                    file_hits.append({
                        "content": content,
                        "location": location,
                        "types": ", ".join(types),
                        "positions": positions,
                        "extra": extra,
                        "extra2": extra2,
                    })

            if ext == '.txt' and not file_hits and full_text:
                types, positions = check_match_5(full_text, v1, v2, v3, v4, v5)
                if types:
                    first_pos = positions[0] if positions else (0, 0, "")
                    line_num = full_text.count("\n", 0, first_pos[0]) + 1
                    start_ctx = max(0, first_pos[0] - 80)
                    end_ctx = min(len(full_text), first_pos[1] + 80)
                    context = full_text[start_ctx:end_ctx].replace("\n", " ")
                    local_positions = []
                    for p in positions:
                        lp = (p[0] - start_ctx, p[1] - start_ctx, p[2])
                        if 0 <= lp[0] and lp[1] <= len(context):
                            local_positions.append(lp)
                    file_hits.append({
                        "content": context,
                        "location": f"Многострочное совпадение (строка ~{line_num})",
                        "types": ", ".join(types),
                        "positions": local_positions,
                        "extra": "",
                        "extra2": "",
                    })

            if file_hits:
                if not _file_hits_pass_filter(file_hits):
                    continue

                header = tk.Frame(result_text, bg=COLORS["result_bg"])
                result_text.window_create(tk.END, window=header)
                result_text.insert(tk.END, "\n", "green")

                tk.Label(header, text=f"📁 Файл: {fname}",
                         font=("Consolas", 10, "bold"),
                         fg="green", bg=COLORS["result_bg"]).pack(side=tk.LEFT)
                tk.Label(header, text=f"   Путь: {filepath}",
                         font=("Consolas", 9),
                         fg="gray", bg=COLORS["result_bg"]).pack(side=tk.LEFT)

                btn_folder = tk.Button(header, text="📁 Папка",
                    command=lambda fp=filepath: show_in_folder(fp),
                    font=("Arial", 9), bg="#2196F3", fg="white")
                btn_folder.pack(side=tk.RIGHT, padx=5)

                btn = tk.Button(header, text="📂 Открыть",
                    command=lambda fp=filepath, fh=file_hits, e=ext:
                        open_matches_in_file(fp, fh, e),
                    font=("Arial", 9), bg="#4CAF50", fg="white")
                btn.pack(side=tk.RIGHT, padx=5)

                result_text.insert(tk.END, "-" * 55 + "\n")

                for hit in file_hits:
                    total_hits += 1
                    result_text.insert(tk.END, "   ✅ Совпадение (")
                    result_text.insert(tk.END, hit["types"], "green_bold")
                    result_text.insert(tk.END, ")\n")
                    result_text.insert(tk.END,
                        f"      Где: {hit['location']}\n")
                    result_text.insert(tk.END, "      Что: ")

                    _insert_highlighted(result_text, hit["content"],
                                        hit["positions"])
                    result_text.insert(tk.END, "\n\n")

                    last_results.append({
                        "source": f"Файл: {fname}",
                        "path": filepath,
                        "location": hit["location"],
                        "types": hit["types"],
                        "content": hit["content"],
                        "hits_positions": hit["positions"],
                        "extra": hit["extra"],
                        "extra2": hit["extra2"],
                        "ext": ext,
                    })
            else:
                result_text.insert(tk.END,
                    f"\n❌ Файл: {fname} — совпадений нет\n", "red")

        if not stop_requested:
            result_text.insert(tk.END, "=" * 60 + "\n")
            result_text.insert(tk.END,
                f"Готово. Всего найдено совпадений: {total_hits}\n", "green")
            if total_hits == 0:
                result_text.insert(tk.END,
                    "Совпадений не найдено. Проверьте параметры поиска.\n")
            app_log(f"Поиск завершён. Найдено совпадений: {total_hits}")

    except Exception as e:
        result_text.insert(tk.END,
            f"\n[Ошибка во время поиска] {str(e)}\n", "red")
        app_log(f"ОШИБКА поиска: {e}")
    finally:
        search_running = False
        stop_requested = False
        search_btn.config(state=tk.NORMAL)
        stop_btn.config(state=tk.DISABLED)
        update_progress(100, "Готово.")
        update_status_bar()


def _file_hits_pass_filter(file_hits):
    want_v1 = filter_res_v1_var.get()
    want_v2 = filter_res_v2_var.get()
    want_v3 = filter_res_v3_var.get()
    want_v4 = filter_res_v4_var.get()
    want_v5 = filter_res_v5_var.get()

    if want_v1 and want_v2 and want_v3 and want_v4 and want_v5:
        return True

    for hit in file_hits:
        t = hit.get("types", "")
        if want_v1 and "Значение 1" in t:
            return True
        if want_v2 and "Значение 2" in t:
            return True
        if want_v3 and "Значение 3" in t:
            return True
        if want_v4 and "Значение 4" in t:
            return True
        if want_v5 and "Значение 5" in t:
            return True
    return False


def _insert_highlighted(text_widget, content, positions):
    if not positions:
        text_widget.insert(tk.END, content)
        return
    last = 0
    for start, end, _t in positions:
        if start > last:
            text_widget.insert(tk.END, content[last:start])
        text_widget.insert(tk.END, content[start:end], "green_bold")
        last = end
    if last < len(content):
        text_widget.insert(tk.END, content[last:])


def open_matches_in_file(filepath, file_hits, ext):
    if ext == '.xlsx':
        first = file_hits[0]
        cell_coord = first.get("extra", "")
        sheet_name = first.get("extra2", "")
        if cell_coord and sheet_name:
            open_excel_at(filepath, sheet_name, cell_coord)
        else:
            open_file_generic(filepath)

    elif ext == '.docx':
        first = file_hits[0]
        location = first.get("location", "")
        open_word_at(filepath, location)

    elif ext == '.txt':
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
        except:
            try:
                with open(filepath, "r", encoding="cp1251") as f:
                    content = f.read()
            except:
                content = ""

        all_positions = []
        for hit in file_hits:
            content_hit = hit.get("content", "")
            if "Многострочное" in hit.get("location", ""):
                for needle in (v1_text_global, v2_text_global,
                               v3_text_global, v4_text_global, v5_text_global):
                    if needle:
                        for m in re.finditer(re.escape(needle), content, re.IGNORECASE):
                            all_positions.append((m.start(), m.end()))
            else:
                idx = content.find(content_hit)
                if idx >= 0:
                    for start, end, _t in hit.get("positions", []):
                        all_positions.append((idx + start, idx + end))

        all_positions = sorted(set(all_positions))
        TxtViewerWindow(root, filepath, all_positions)


def start_search():
    global search_running, stop_requested
    global v1_text_global, v2_text_global, v3_text_global
    global v4_text_global, v5_text_global

    if search_running:
        messagebox.showinfo("Поиск уже идёт",
            "Дождитесь окончания или остановите поиск.")
        return

    v1 = v1_combo.get().strip()
    v2 = v2_combo.get().strip()
    v3 = v3_combo.get().strip()
    v4 = v4_combo.get().strip()
    v5 = v5_combo.get().strip()

    v1_text_global = v1
    v2_text_global = v2
    v3_text_global = v3
    v4_text_global = v4
    v5_text_global = v5

    if not selected_files:
        messagebox.showwarning("Внимание",
            "Добавьте хотя бы один файл для поиска.")
        return

    if not v1 and not v2 and not v3 and not v4 and not v5:
        messagebox.showwarning("Внимание",
            "Заполните хотя бы одно поле (Значение 1-5).")
        return

    add_to_history("v1", v1)
    add_to_history("v2", v2)
    add_to_history("v3", v3)
    add_to_history("v4", v4)
    add_to_history("v5", v5)
    reload_history_combos()

    filter_xlsx = bool(filter_xlsx_var.get())
    filter_docx = bool(filter_docx_var.get())
    filter_txt = bool(filter_txt_var.get())

    result_text.delete(1.0, tk.END)
    update_progress(0, "Поиск...")
    update_status_bar("Поиск запущен...")
    search_btn.config(state=tk.DISABLED)
    stop_btn.config(state=tk.NORMAL)

    search_running = True
    stop_requested = False

    t = threading.Thread(
        target=search_thread_target,
        args=(v1, v2, v3, v4, v5, filter_xlsx, filter_docx, filter_txt),
        daemon=True
    )
    t.start()


def stop_search():
    global stop_requested
    if search_running:
        stop_requested = True
        update_status_bar("Остановка поиска...")
        app_log("Пользователь остановил поиск")


# ============================================================
#  ИНТЕРФЕЙС
# ============================================================
root = tk.Tk()
root.title("VS Search Engine v1.0")
root.geometry("980x1000")

set_colors("light")

BTN_SEARCH = ("#4CAF50", "#ffffff")
BTN_STOP = ("#f44336", "#ffffff")
BTN_THEME = ("#FFC107", "#000000")
BTN_HELP = ("#2196F3", "#ffffff")
BTN_EXPORT = ("#2196F3", "#ffffff")

# ---------- ФАЙЛЫ ----------
frame_files = tk.LabelFrame(root, text="ФАЙЛЫ ДЛЯ ПОИСКА", padx=5, pady=5)
frame_files.pack(fill=tk.BOTH, expand=False, padx=10, pady=6)

list_frame = tk.Frame(frame_files)
list_frame.pack(fill=tk.BOTH, expand=True)

scrollbar = tk.Scrollbar(list_frame)
scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

listbox = tk.Listbox(list_frame, yscrollcommand=scrollbar.set,
                     height=5, selectmode=tk.EXTENDED)
listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
scrollbar.config(command=listbox.yview)

btn_frame = tk.Frame(frame_files)
btn_frame.pack(fill=tk.X, pady=5)

tk.Button(btn_frame, text="➕ Добавить файлы",
          command=add_files, width=20).pack(side=tk.LEFT, padx=2)
tk.Button(btn_frame, text="➖ Удалить выбранный",
          command=remove_selected, width=22).pack(side=tk.LEFT, padx=2)
tk.Button(btn_frame, text="🗑 Очистить список",
          command=clear_list, width=20).pack(side=tk.LEFT, padx=2)

filter_frame = tk.Frame(frame_files)
filter_frame.pack(fill=tk.X, pady=3)
tk.Label(filter_frame, text="Искать в форматах:").pack(side=tk.LEFT, padx=(0, 10))
filter_xlsx_var = tk.BooleanVar(value=True)
filter_docx_var = tk.BooleanVar(value=True)
filter_txt_var = tk.BooleanVar(value=True)
tk.Checkbutton(filter_frame, text="Excel (.xlsx)",
               variable=filter_xlsx_var).pack(side=tk.LEFT)
tk.Checkbutton(filter_frame, text="Word (.docx)",
               variable=filter_docx_var).pack(side=tk.LEFT)
tk.Checkbutton(filter_frame, text="Текст (.txt)",
               variable=filter_txt_var).pack(side=tk.LEFT)

# ---------- ПАРАМЕТРЫ (5 ЗНАЧЕНИЙ) ----------
frame_params = tk.LabelFrame(root, text="ПАРАМЕТРЫ ПОИСКА (5 значений)",
                              padx=5, pady=5)
frame_params.pack(fill=tk.X, expand=False, padx=10, pady=6)

r1 = tk.Frame(frame_params)
r1.pack(fill=tk.X, pady=3)
tk.Label(r1, text="Значение 1:", width=14, anchor="w").pack(side=tk.LEFT)
v1_var = tk.StringVar()
v1_combo = ttk.Combobox(r1, textvariable=v1_var, values=[])
v1_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

r2 = tk.Frame(frame_params)
r2.pack(fill=tk.X, pady=3)
tk.Label(r2, text="Значение 2:", width=14, anchor="w").pack(side=tk.LEFT)
v2_var = tk.StringVar()
v2_combo = ttk.Combobox(r2, textvariable=v2_var, values=[])
v2_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

r3 = tk.Frame(frame_params)
r3.pack(fill=tk.X, pady=3)
tk.Label(r3, text="Значение 3:", width=14, anchor="w").pack(side=tk.LEFT)
v3_var = tk.StringVar()
v3_combo = ttk.Combobox(r3, textvariable=v3_var, values=[])
v3_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

r4 = tk.Frame(frame_params)
r4.pack(fill=tk.X, pady=3)
tk.Label(r4, text="Значение 4:", width=14, anchor="w").pack(side=tk.LEFT)
v4_var = tk.StringVar()
v4_combo = ttk.Combobox(r4, textvariable=v4_var, values=[])
v4_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

r5 = tk.Frame(frame_params)
r5.pack(fill=tk.X, pady=3)
tk.Label(r5, text="Значение 5:", width=14, anchor="w").pack(side=tk.LEFT)
v5_var = tk.StringVar()
v5_combo = ttk.Combobox(r5, textvariable=v5_var, values=[])
v5_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

tk.Button(frame_params, text="🗑 Очистить историю",
          command=clear_history).pack(anchor="w", pady=4)

# ---------- КНОПКИ ДЕЙСТВИЙ ----------
btn_actions = tk.Frame(root)
btn_actions.pack(fill=tk.X, padx=10, pady=6)

search_btn = tk.Button(btn_actions, text="🔍 ПОИСК",
                       command=start_search,
                       font=("Arial", 12, "bold"), height=2,
                       bg=BTN_SEARCH[0], fg=BTN_SEARCH[1],
                       activebackground="#3d8b40")
search_btn.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

stop_btn = tk.Button(btn_actions, text="⏹ Остановить",
                     command=stop_search,
                     font=("Arial", 11), height=2, width=14,
                     state=tk.DISABLED,
                     bg=BTN_STOP[0], fg=BTN_STOP[1],
                     activebackground="#c62828")
stop_btn.pack(side=tk.LEFT, padx=(5, 5))

tk.Button(btn_actions, text="🎨 Тема", command=toggle_theme,
          font=("Arial", 11), height=2, width=10,
          bg=BTN_THEME[0], fg=BTN_THEME[1],
          activebackground="#d39e00").pack(side=tk.LEFT, padx=(0, 5))

tk.Button(btn_actions, text="❓ Руководство", command=show_guide,
          font=("Arial", 11), height=2, width=14,
          bg=BTN_HELP[0], fg=BTN_HELP[1],
          activebackground="#1976D2").pack(side=tk.LEFT, padx=(0, 5))

btn_export = tk.Frame(root)
btn_export.pack(fill=tk.X, padx=10, pady=(0, 6))

tk.Button(btn_export, text="💾 Экспорт TXT",
          command=export_txt, width=22,
          bg=BTN_EXPORT[0], fg=BTN_EXPORT[1],
          activebackground="#1976D2").pack(side=tk.LEFT, padx=2)
tk.Button(btn_export, text="📊 Экспорт XLSX",
          command=export_xlsx, width=22,
          bg=BTN_EXPORT[0], fg=BTN_EXPORT[1],
          activebackground="#1976D2").pack(side=tk.LEFT, padx=2)

# ---------- ПРОГРЕСС ----------
progress_frame = tk.Frame(root)
progress_frame.pack(fill=tk.X, padx=10, pady=4)

progress_bar = ttk.Progressbar(progress_frame, orient="horizontal",
                                mode="determinate", length=600)
progress_bar.pack(fill=tk.X, side=tk.LEFT, expand=True)

progress_label = tk.Label(progress_frame, text="Готов 0%", width=34)
progress_label.pack(side=tk.LEFT, padx=(8, 0))

# ---------- ФИЛЬТР РЕЗУЛЬТАТОВ ----------
frame_filter = tk.LabelFrame(root, text="ФИЛЬТР РЕЗУЛЬТАТОВ",
                              padx=5, pady=3)
frame_filter.pack(fill=tk.X, expand=False, padx=10, pady=(0, 6))

f1 = tk.Frame(frame_filter)
f1.pack(fill=tk.X, pady=2)
tk.Label(f1, text="Показать:", width=12, anchor="w").pack(side=tk.LEFT)
filter_res_v1_var = tk.BooleanVar(value=True)
filter_res_v2_var = tk.BooleanVar(value=True)
filter_res_v3_var = tk.BooleanVar(value=True)
filter_res_v4_var = tk.BooleanVar(value=True)
filter_res_v5_var = tk.BooleanVar(value=True)
tk.Checkbutton(f1, text="Значение 1", variable=filter_res_v1_var).pack(side=tk.LEFT, padx=5)
tk.Checkbutton(f1, text="Значение 2", variable=filter_res_v2_var).pack(side=tk.LEFT, padx=5)
tk.Checkbutton(f1, text="Значение 3", variable=filter_res_v3_var).pack(side=tk.LEFT, padx=5)
tk.Checkbutton(f1, text="Значение 4", variable=filter_res_v4_var).pack(side=tk.LEFT, padx=5)
tk.Checkbutton(f1, text="Значение 5", variable=filter_res_v5_var).pack(side=tk.LEFT, padx=5)

# ---------- РЕЗУЛЬТАТЫ ----------
tk.Label(root, text="РЕЗУЛЬТАТЫ:", anchor="w").pack(fill=tk.X, padx=10)

result_text = scrolledtext.ScrolledText(root, wrap=tk.WORD,
                                        font=("Consolas", 10))
result_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

result_text.tag_config("green", foreground="green")
result_text.tag_config("red", foreground="red")
result_text.tag_config("yellow", foreground="#B8860B")
result_text.tag_config("green_bold",
    foreground="#006400",
    font=("Consolas", 10, "bold"),
    background="#C8E6C9")

# ---------- СТАТУС-БАР ----------
status_frame = tk.Frame(root, relief=tk.SUNKEN, bd=1)
status_frame.pack(fill=tk.X, side=tk.BOTTOM)

status_bar_label = tk.Label(status_frame, text="Готов",
                            anchor="w", font=("Arial", 9))
status_bar_label.pack(fill=tk.X, padx=6, pady=2)

# ---------- ЗАГРУЗКА НАСТРОЕК ----------
try:
    cfg = load_config()
    filter_xlsx_var.set(cfg.get("filter_xlsx", True))
    filter_docx_var.set(cfg.get("filter_docx", True))
    filter_txt_var.set(cfg.get("filter_txt", True))
    current_theme = cfg.get("theme", "light")
    filter_res_v1_var.set(cfg.get("filter_res_v1", True))
    filter_res_v2_var.set(cfg.get("filter_res_v2", True))
    filter_res_v3_var.set(cfg.get("filter_res_v3", True))
    filter_res_v4_var.set(cfg.get("filter_res_v4", True))
    filter_res_v5_var.set(cfg.get("filter_res_v5", True))
except Exception as e:
    print("Ошибка загрузки настроек:", e)

apply_current_theme()
restore_pinned_files()
reload_history_combos()

# ---------- КОНТЕКСТНОЕ МЕНЮ ----------
def apply_context_menu_recursive(parent):
    for child in parent.winfo_children():
        if isinstance(child, (tk.Entry, ttk.Entry, ttk.Combobox)):
            add_context_menu(child)
        apply_context_menu_recursive(child)


apply_context_menu_recursive(root)
add_result_context_menu(result_text)
add_context_menu(listbox)

# ---------- ГОРЯЧИЕ КЛАВИШИ ----------
def clear_results():
    result_text.delete(1.0, tk.END)
    update_status_bar()


root.bind("<Control-Return>", lambda e: start_search())
root.bind("<Control-o>", lambda e: add_files())
root.bind("<Control-O>", lambda e: add_files())
root.bind("<Control-l>", lambda e: clear_list())
root.bind("<Control-L>", lambda e: clear_list())
root.bind("<Control-e>", lambda e: export_txt())
root.bind("<Control-E>", lambda e: export_txt())
root.bind("<Control-Shift-E>", lambda e: export_xlsx())
root.bind("<Control-Shift-e>", lambda e: export_xlsx())
root.bind("<Control-d>", lambda e: toggle_theme())
root.bind("<Control-D>", lambda e: toggle_theme())
root.bind("<Control-h>", lambda e: clear_history())
root.bind("<Control-H>", lambda e: clear_history())
root.bind("<F1>", lambda e: show_guide())
root.bind("<F5>", lambda e: clear_results())
root.bind("<Escape>", lambda e: stop_search())

root.protocol("WM_DELETE_WINDOW", on_close)

app_log_start()
update_status_bar()

root.mainloop()